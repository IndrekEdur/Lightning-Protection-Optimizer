"""Signed calculation exports: XLSX material list and PDF calculation report."""
import hashlib
import hmac
import io
import json
import secrets
from xml.sax.saxutils import escape
from .models import Rod

_SECRET=secrets.token_bytes(32)


def token(project,result):
    body={key:value for key,value in result.items() if key!='report_token'}
    normalized=project.model_copy(update={'rods':[Rod(**r) for r in result['rods']]})
    def canonical(value):
        if isinstance(value,float) and value.is_integer():return int(value)
        if isinstance(value,dict):return {k:canonical(v) for k,v in value.items()}
        if isinstance(value,(list,tuple)):return [canonical(v) for v in value]
        return value
    data=json.dumps(canonical({'project':normalized.model_dump(),'result':body}),sort_keys=True,separators=(',',':'),allow_nan=False).encode()
    return hmac.new(_SECRET,data,hashlib.sha256).hexdigest()


def check_token(project,result):
    supplied=result.get('report_token','')
    if not isinstance(supplied,str) or not hmac.compare_digest(supplied,token(project,result)):
        raise ValueError('Calculation changed or server restarted. Recalculate before exporting.')


def facts(p,r):
    verification=r.get('verification',{})
    ifc_facts=[('IFC source',p.ifc_metadata.filename),('IFC schema',p.ifc_metadata.schema_name),('Source unit to metre',p.ifc_metadata.length_unit_scale_to_m),('IFC source SHA256',p.ifc_metadata.source_sha256),('Coordinate reference',p.ifc_metadata.coordinate_reference),('Ground reference confirmed',str(p.reference_surface_confirmed)),('IFC classifications reviewed',str(p.ifc_review_confirmed)),('Excluded missing IFC elements',', '.join(p.acknowledged_ifc_omissions)),('IFC tessellation/diagnostics',json.dumps(p.ifc_metadata.diagnostics))] if p.ifc_metadata else []
    return ifc_facts+[('Project',p.name),('Method',r.get('method',p.method)),('LPS',p.lps),('Sphere radius (m)',r.get('radius','n/a')),('Standard edition',p.edition),('Units',p.units),('Geometry status',r.get('verification_status','demo only')),('Continuous geometry certificate',str(r.get('geometry_verified',False))),('IEC/EN compliance established','No'),('Ground reference (m)',p.ground_z),('Geometry tolerance (m)',p.sphere_tolerance),('Sample spacing (m)',p.sample_spacing),('Surface certificate max depth',verification.get('max_depth','n/a')),('Certificate nodes',verification.get('nodes','n/a')),('Certificate unresolved patches',verification.get('unresolved_patches','n/a')),('Certificate proof',verification.get('proof','n/a')),('Rod candidate spacing (m)',p.spacing),('Mounting surfaces',', '.join(p.candidate_surfaces)),('Mount slope normal z minimum',p.min_mount_normal_z),('Objective',p.objective),('Solver status',r['solver_status']),('MILP gap (witness model)',r.get('mip_gap')),('Search method',r.get('optimization_method','finite grid synthetic set cover')),('Total rod cost',r['total_cost']),('Rod count',len(r['rods'])),('Duration (seconds)',r.get('duration_seconds','n/a'))]


def schedule(p,r):
    return [[i+1,p.catalogue[rod['type_index']].name,rod['x'],rod['y'],p.roof_height if rod.get('base_z') is None else rod['base_z'],p.catalogue[rod['type_index']].height,p.catalogue[rod['type_index']].price] for i,rod in enumerate(r['rods'])]


def xlsx(p,r):
    from openpyxl import Workbook
    book=Workbook();summary=book.active;summary.title='Calculation'
    for row in facts(p,r):summary.append(row)
    rods=book.create_sheet('Rods');rods.append(['ID','Type','x (m)','y (m)','Base z (m)','Height (m)','Price'])
    for row in schedule(p,r):rods.append(row)
    objects=book.create_sheet('Object status');objects.append(['Object','Status'])
    for name,status in r.get('verification',{}).get('object_status',{}).items():objects.append([name,status])
    contacts=book.create_sheet('Unprotected contacts');contacts.append(['Object','x','y','z','Sphere cx','Sphere cy','Sphere cz','Clearance'])
    for pt in r.get('verification',{}).get('points',[]):
        if pt['status']=='unprotected':contacts.append([pt['object'],*pt['point'],*pt['witness']['center'],pt['witness']['clearance']])
    checks=book.create_sheet('Engineering checks')
    for text in r['checks']:checks.append([text])
    for sheet in book:
        sheet.freeze_panes='A2'
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(75,max(12,max(len(str(c.value or '')) for c in col)+2))
    buffer=io.BytesIO();book.save(buffer);return buffer.getvalue()


def pdf(p,r):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    buffer=io.BytesIO();style=getSampleStyleSheet();story=[]
    def text(value):return Paragraph(escape(str(value)),style['BodyText'])
    story.extend([Paragraph('Lightning Protection Optimizer',style['Title']),text('Calculation report - geometry verification, not an IEC/EN compliance certificate'),Spacer(1,12)])
    for key,value in facts(p,r):story.append(text(f'{key}: {value}'))
    story.extend([Spacer(1,12),Paragraph('Rod schedule (metres)',style['Heading2'])])
    table=Table([['ID','Type','x','y','Base z','Height','Price']]+[[text(f'{v:.3f}' if isinstance(v,float) else v) for v in row] for row in schedule(p,r)],repeatRows=1,hAlign='LEFT',colWidths=[30,45,65,65,65,65,70])
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.lightgrey),('GRID',(0,0),(-1,-1),.3,colors.grey)]));story.append(table)
    story.append(Paragraph('Object protection status',style['Heading2']))
    for name,status in r.get('verification',{}).get('object_status',{}).items():story.append(text(f'{name}: {status}'))
    story.append(Paragraph('Detected unprotected contact locations',style['Heading2']))
    for pt in r.get('verification',{}).get('points',[]):
        if pt['status']=='unprotected':story.append(text(f"{pt['object']}: target {pt['point']}; sphere center {pt['witness']['center']}; clearance {pt['witness']['clearance']} m"))
    story.append(Paragraph('Outstanding engineering checks and assumptions',style['Heading2']))
    for check in r['checks']:story.append(text(check))
    SimpleDocTemplate(buffer).build(story);return buffer.getvalue()
