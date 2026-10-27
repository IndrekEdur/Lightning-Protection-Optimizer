import io
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from lpo.api import app


def test_real_pdf_xlsx_and_tamper_detection():
    client=TestClient(app)
    project={'method':'rolling_sphere','objective':'manual','width':2,'depth':2,'rods':[{'x':1,'y':1,'type_index':1}],'sphere_refinement':0}
    response=client.post('/api/optimize',json=project)
    assert response.status_code==200
    result=response.json()
    assert result['geometry_verified']
    payload={'project':project,'result':result}
    excel=client.post('/api/report/xlsx',json=payload)
    assert excel.status_code==200
    book=load_workbook(io.BytesIO(excel.content))
    assert book['Rods'].max_row==2
    assert book['Rods']['F2'].value==5
    assert book['Object status']['B2'].value=='protected'
    pdf=client.post('/api/report/pdf',json=payload)
    assert pdf.status_code==200
    assert pdf.content.startswith(b'%PDF')
    result['total_cost']=0
    assert client.post('/api/report/pdf',json=payload).status_code==422


def test_changed_geometry_cannot_reuse_report_certificate():
    client=TestClient(app)
    project={'method':'rolling_sphere','objective':'manual','width':2,'depth':2}
    result=client.post('/api/optimize',json=project).json()
    project['width']=20
    assert client.post('/api/report/xlsx',json={'project':project,'result':result}).status_code==422


def test_browser_numeric_roundtrip_preserves_report_signature():
    client=TestClient(app)
    project={'method':'rolling_sphere','objective':'manual','width':2,'depth':2}
    result=client.post('/api/optimize',json=project).json()
    def javascript_numbers(value):
        if isinstance(value,float) and value.is_integer():return int(value)
        if isinstance(value,dict):return {k:javascript_numbers(v) for k,v in value.items()}
        if isinstance(value,list):return [javascript_numbers(v) for v in value]
        return value
    result=javascript_numbers(result)
    assert client.post('/api/report/xlsx',json={'project':project,'result':result}).status_code==200
